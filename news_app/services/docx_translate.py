"""原文と和訳の Word（.docx）出力。"""
from __future__ import annotations

import io

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from news_app.services.openai_translate import split_script_units
from news_app.services.storage import _normalize_script_ja_pairs

FONT_NAME = "Yu Gothic"
HEADER_FILL = "F0FDFA"
HEADER_EN_COLOR = RGBColor(0x64, 0x74, 0x8B)
HEADER_JA_COLOR = RGBColor(0x92, 0x40, 0x0E)
BODY_EN_COLOR = RGBColor(0x1E, 0x29, 0x3B)
BODY_JA_COLOR = RGBColor(0x33, 0x41, 0x55)


def translation_rows(script: str, translation: str, pairs) -> list[dict]:
    """画面表示と同じルールで原文・和訳の行を揃える。"""
    normalized = _normalize_script_ja_pairs(pairs)
    if normalized:
        return normalized
    units = split_script_units(script)
    ja_units = split_script_units(translation)
    if not units:
        if not str(translation or "").strip():
            return []
        return [{"en": "", "ja": str(translation).strip()}]
    return [
        {
            "en": en,
            "ja": ja_units[index] if len(ja_units) == len(units) else (translation if index == 0 else ""),
        }
        for index, en in enumerate(units)
    ]


def _set_run_font(run, *, size_pt: float, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size_pt)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), FONT_NAME)
    r_fonts.set(qn("w:hAnsi"), FONT_NAME)
    r_fonts.set(qn("w:eastAsia"), FONT_NAME)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _set_cell_margins(cell, *, top=60, bottom=60, left=80, right=80) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def _fill_cell(cell, text: str, *, header: bool = False, japanese: bool = False) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    _set_cell_margins(cell)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15
    display = str(text or "").strip() or ("（和訳なし）" if japanese else "（原文なし）")
    run = paragraph.add_run(display)
    if header:
        _set_run_font(run, size_pt=10, bold=True, color=HEADER_JA_COLOR if japanese else HEADER_EN_COLOR)
    else:
        _set_run_font(run, size_pt=11, color=BODY_JA_COLOR if japanese else BODY_EN_COLOR)


def build_script_translation_docx(*, title: str = "", pairs: list[dict]) -> bytes:
    """原文と和訳の2列表を含む .docx バイナリを返す。"""
    rows = _normalize_script_ja_pairs(pairs)
    if not rows:
        raise ValueError("原文と和訳がありません。")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)

    heading = doc.add_paragraph()
    heading.paragraph_format.space_after = Pt(2)
    heading_run = heading.add_run("原文と和訳")
    _set_run_font(heading_run, size_pt=16, bold=True, color=BODY_EN_COLOR)

    cleaned_title = str(title or "").strip()
    if cleaned_title:
        subtitle = doc.add_paragraph()
        subtitle.paragraph_format.space_before = Pt(0)
        subtitle.paragraph_format.space_after = Pt(10)
        subtitle_run = subtitle.add_run(cleaned_title)
        _set_run_font(subtitle_run, size_pt=11, color=HEADER_EN_COLOR)
    else:
        heading.paragraph_format.space_after = Pt(10)

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    usable_width = int(section.page_width - section.left_margin - section.right_margin)
    col_width = usable_width // 2
    for column in table.columns:
        column.width = col_width

    header_cells = table.rows[0].cells
    _set_cell_shading(header_cells[0], HEADER_FILL)
    _set_cell_shading(header_cells[1], HEADER_FILL)
    _fill_cell(header_cells[0], "原文", header=True, japanese=False)
    _fill_cell(header_cells[1], "和訳", header=True, japanese=True)
    header_cells[0].width = col_width
    header_cells[1].width = col_width

    for row in rows:
        cells = table.add_row().cells
        cells[0].width = col_width
        cells[1].width = col_width
        _fill_cell(cells[0], row.get("en") or "", japanese=False)
        _fill_cell(cells[1], row.get("ja") or "", japanese=True)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
